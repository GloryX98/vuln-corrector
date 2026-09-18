#!/usr/bin/env python3
"""
vuln_corrector - Correct Vulnerability Rating and CVSS in a Nessus/Tenable vuln sheet.

For every finding in the sheet it looks the vulnerability up on Tenable (by CVE, with
an NVD fallback, and a best-effort Tenable plugin-name search when no CVE is present),
derives the authoritative CVSS v3 base score + risk factor, and rewrites ONLY the
vulnerability-rating and CVSS columns.

Rules (match the manual review):
  * Rating  = Tenable/NVD risk factor (Low / Medium / High / Critical)
  * CVSS     = base score rounded to an integer, then CLAMPED into the rating band:
                 Low 2-3 | Medium 4-6 | High 7-8 | Critical 9-10
  * When a finding maps to several CVEs, the worst (highest) score wins - same as Nessus.
  * Every other cell, column, row order and the header are left exactly as they were.
  * The input file is never modified; a corrected copy is written next to it
    (or to -o / --output). Use --inplace only if you really want to overwrite.

Needs an internet connection (queries tenable.com / services.nvd.nist.gov).
Pure standard library for HTTP - runs on Windows, macOS and Linux with a stock
Python 3 (openpyxl is only needed for .xlsx input).

Usage (Windows .exe / macOS-Linux binary / from source):
    vuln_corrector      findings.csv
    vuln_corrector      findings.xlsx -o fixed.xlsx --report changes.csv
    python3 vuln_corrector.py  findings.csv        # run from source, any OS
    vuln_corrector                                 # no args -> prompts for a path
"""
from __future__ import annotations
import argparse
import csv
import math
import os
import re
import sys
import time
import json
import platform
import urllib.request
import urllib.error
from urllib.parse import quote_plus

# Cross-platform: no third-party HTTP dependency. Uses the standard library only,
# so the source runs as-is on Windows, macOS and Linux with a stock Python 3.
HEADERS = {
    "User-Agent": "Mozilla/5.0 (%s) vuln-corrector/1.1" % platform.system(),
    "Accept": "text/html,application/json,*/*",
}
TIMEOUT = 25


def http_get(url: str, timeout: int = TIMEOUT):
    """GET a URL with the stdlib. Returns (status_code|None, body_text)."""
    req = urllib.request.Request(url, headers=HEADERS)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            charset = resp.headers.get_content_charset() or "utf-8"
            return getattr(resp, "status", 200), resp.read().decode(charset, "replace")
    except urllib.error.HTTPError as e:            # 4xx / 5xx
        return e.code, ""
    except (urllib.error.URLError, TimeoutError, OSError, ValueError):
        return None, ""

# --- column name auto-detection -------------------------------------------------
TITLE_KEYS  = ["vulnerability_title", "vulnerability_name", "plugin_name", "title", "name", "vulnerability"]
RATING_KEYS = ["vulnerability_rating", "risk_factor", "risk", "severity", "rating"]
CVSS_KEYS   = ["cvss", "cvss_score", "cvss3", "cvssv3", "cvss_base", "base_score", "score"]
CVE_KEYS    = ["cve", "cves", "cve_id", "cve_ids"]

CVE_RE = re.compile(r"CVE-\d{4}-\d{4,7}", re.IGNORECASE)
V3_VECTOR_RE = re.compile(r"CVSS:3\.[01]/(?:[A-Z]+:[A-Z]+/?)+")

# --- CVSS v3.1 base-score calculator (deterministic, from the vector) -----------
_W = {
    "AV": {"N": 0.85, "A": 0.62, "L": 0.55, "P": 0.20},
    "AC": {"L": 0.77, "H": 0.44},
    "PR_U": {"N": 0.85, "L": 0.62, "H": 0.27},   # scope Unchanged
    "PR_C": {"N": 0.85, "L": 0.68, "H": 0.50},   # scope Changed
    "UI": {"N": 0.85, "R": 0.62},
    "CIA": {"N": 0.0, "L": 0.22, "H": 0.56},
}

def _roundup(x: float) -> float:
    # official CVSS Roundup: smallest 1-decimal value >= x (with integer-scaled math)
    i = int(round(x * 100000))
    if i % 10000 == 0:
        return i / 100000.0
    return (math.floor(i / 10000) + 1) / 10.0

def base_score_from_vector(vector: str):
    """Return (base_score, severity) for a CVSS:3.x vector string, or (None, None)."""
    try:
        parts = dict(p.split(":") for p in vector.split("/") if ":" in p and not p.startswith("CVSS"))
        av = _W["AV"][parts["AV"]]
        ac = _W["AC"][parts["AC"]]
        ui = _W["UI"][parts["UI"]]
        scope_changed = parts["S"] == "C"
        pr = (_W["PR_C"] if scope_changed else _W["PR_U"])[parts["PR"]]
        c = _W["CIA"][parts["C"]]; i_ = _W["CIA"][parts["I"]]; a = _W["CIA"][parts["A"]]
    except (KeyError, ValueError):
        return None, None
    isc_base = 1 - (1 - c) * (1 - i_) * (1 - a)
    if scope_changed:
        impact = 7.52 * (isc_base - 0.029) - 3.25 * (isc_base - 0.02) ** 15
    else:
        impact = 6.42 * isc_base
    expl = 8.22 * av * ac * pr * ui
    if impact <= 0:
        base = 0.0
    elif scope_changed:
        base = _roundup(min(1.08 * (impact + expl), 10))
    else:
        base = _roundup(min(impact + expl, 10))
    return base, severity_from_score(base)

def severity_from_score(score: float) -> str:
    if score is None:
        return "Unknown"
    if score == 0:      return "None"
    if score < 4.0:     return "Low"
    if score < 7.0:     return "Medium"
    if score < 9.0:     return "High"
    return "Critical"

# --- integer CVSS clamped into the user's rating band ---------------------------
BANDS = {"Low": (2, 3), "Medium": (4, 6), "High": (7, 8), "Critical": (9, 10)}

def clamped_integer(score: float, severity: str) -> int:
    n = int(math.floor(score + 0.5))          # round half up
    lo, hi = BANDS.get(severity, (n, n))
    return max(lo, min(hi, n))

# --- lookups --------------------------------------------------------------------
_cache: dict[str, tuple] = {}

def lookup_cve_tenable(cve: str):
    """Return (base_score, severity) from the Tenable CVE page, else (None, None)."""
    url = f"https://www.tenable.com/cve/{cve.upper()}"
    status, text = http_get(url)
    if status != 200 or not text:
        return None, None
    best = None
    for v in V3_VECTOR_RE.findall(text):
        s, sev = base_score_from_vector(v)
        if s is not None and (best is None or s > best[0]):
            best = (s, sev)
    return best if best else (None, None)

def lookup_cve_nvd(cve: str):
    """Fallback: NVD 2.0 API (structured JSON)."""
    url = f"https://services.nvd.nist.gov/rest/json/cves/2.0?cveId={cve.upper()}"
    for attempt in range(3):
        status, text = http_get(url)
        if status in (403, 429):
            time.sleep(6)                     # NVD rate limit without an API key
            continue
        if status != 200 or not text:
            return None, None
        try:
            data = json.loads(text)
            metrics = (data.get("vulnerabilities") or [{}])[0].get("cve", {}).get("metrics", {})
            for key in ("cvssMetricV31", "cvssMetricV30"):
                if metrics.get(key):
                    cd = metrics[key][0]["cvssData"]
                    score = float(cd["baseScore"])
                    return score, (cd.get("baseSeverity", "").title() or severity_from_score(score))
            return None, None
        except (ValueError, KeyError, IndexError):
            time.sleep(3)
    return None, None

def lookup_cve(cve: str):
    if cve in _cache:
        return _cache[cve]
    res = lookup_cve_tenable(cve)
    if res[0] is None:
        res = lookup_cve_nvd(cve)
    _cache[cve] = res
    return res

def lookup_plugin_by_name(name: str):
    """Best-effort: search Tenable plugins by title, read the first plugin page."""
    key = "PLUGIN::" + name
    if key in _cache:
        return _cache[key]
    result = (None, None)
    surl = "https://www.tenable.com/plugins/search?q=" + quote_plus(f'"{name}"')
    status, text = http_get(surl)
    if status == 200 and text:
        m = re.search(r"/plugins/nessus/(\d+)", text)
        if m:
            pstatus, ptext = http_get(f"https://www.tenable.com/plugins/nessus/{m.group(1)}")
            if pstatus == 200 and ptext:
                best = None
                for v in V3_VECTOR_RE.findall(ptext):
                    s, sev = base_score_from_vector(v)
                    if s is not None and (best is None or s > best[0]):
                        best = (s, sev)
                if best:
                    result = best
                else:
                    rf = re.search(r"Risk Factor[:<>\s\"/a-z]*?(Critical|High|Medium|Low)", ptext, re.I)
                    if rf:
                        sev = rf.group(1).title()
                        lo, hi = BANDS[sev]
                        result = (float(hi), sev)  # no numeric base -> band top as a placeholder
    _cache[key] = result
    return result

def resolve_finding(title: str, cve_field: str):
    """Return (base_score, severity, source) for one finding.

    A Nessus plugin's severity is that of its WORST CVE, so we take the maximum of:
      * the worst CVE listed on the row (Tenable CVE page, NVD fallback), and
      * the Tenable plugin looked up by its exact title
    The row may under-list CVEs, so the plugin-name lookup guards against that.
    """
    candidates = []  # (base, severity, source)

    cves = CVE_RE.findall(cve_field or "") or CVE_RE.findall(title or "")
    if cves:
        best = None; src = ""
        for cve in dict.fromkeys(c.upper() for c in cves):   # dedupe, keep order
            s, sev = lookup_cve(cve)
            if s is not None and (best is None or s > best[0]):
                best = (s, sev); src = cve
        if best:
            candidates.append((best[0], best[1], f"CVE {src}"))

    s, sev = lookup_plugin_by_name(title)
    if s is not None:
        candidates.append((s, sev, "Tenable plugin"))

    if not candidates:
        return None, None, "not found"
    return max(candidates, key=lambda c: c[0])

# --- sheet I/O (CSV + XLSX), preserving everything else -------------------------
def detect_col(fieldnames, keys):
    lower = {f.lower().strip(): f for f in fieldnames}
    for k in keys:
        if k in lower:
            return lower[k]
    for k in keys:                                   # loose contains-match
        for lf, orig in lower.items():
            if k in lf:
                return orig
    return None

def process_rows(rows, fieldnames):
    tcol = detect_col(fieldnames, TITLE_KEYS)
    rcol = detect_col(fieldnames, RATING_KEYS)
    ccol = detect_col(fieldnames, CVSS_KEYS)
    vcol = detect_col(fieldnames, CVE_KEYS)
    if not (tcol and rcol and ccol):
        raise SystemExit(f"Could not find required columns. "
                         f"title={tcol!r} rating={rcol!r} cvss={ccol!r}. "
                         f"Use --title-col/--rating-col/--cvss-col to set them.")
    print(f"Columns -> title:{tcol!r}  rating:{rcol!r}  cvss:{ccol!r}  cve:{vcol!r}\n")

    # resolve each distinct title once
    distinct = {}
    for row in rows:
        distinct.setdefault(row.get(tcol, "").strip(), row.get(vcol, "") if vcol else "")

    resolved = {}
    for i, (title, cve_field) in enumerate(distinct.items(), 1):
        if not title:
            continue
        base, sev, src = resolve_finding(title, cve_field)
        resolved[title] = (base, sev, src)
        if base is None:
            print(f"[{i}/{len(distinct)}] ?  {title[:60]:60}  -> NOT FOUND ({src})")
        else:
            print(f"[{i}/{len(distinct)}] +  {title[:60]:60}  -> {sev} {base} ({src})")
        time.sleep(0.4)                              # be polite to the servers

    changes = []
    for row in rows:
        title = row.get(tcol, "").strip()
        base, sev, src = resolved.get(title, (None, None, "not found"))
        if base is None or sev in ("None", "Unknown"):
            continue
        new_rating = sev
        new_cvss = str(clamped_integer(base, sev))
        old_rating = row.get(rcol, "")
        old_cvss = row.get(ccol, "")
        if str(old_rating) != new_rating or str(old_cvss) != new_cvss:
            changes.append((title, old_rating, new_rating, old_cvss, new_cvss, src))
            row[rcol] = new_rating
            row[ccol] = new_cvss
    return rows, changes, (rcol, ccol)

def read_csv(path):
    with open(path, "r", newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        return list(reader), reader.fieldnames

def write_csv(path, rows, fieldnames):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        writer.writerows(rows)

def read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = {}
        for i in range(len(header)):
            val = r[i] if i < len(r) else None
            row[header[i]] = "" if val is None else val
        rows.append(row)
    return rows, header, wb, ws

def write_xlsx_inplace(wb, ws, header, rows, cols):
    rcol, ccol = cols
    ri = header.index(rcol); ci = header.index(ccol)
    for n, row in enumerate(rows, start=2):
        ws.cell(row=n, column=ri + 1).value = row[rcol]
        ws.cell(row=n, column=ci + 1).value = row[ccol]
    return wb

def main():
    ap = argparse.ArgumentParser(description="Correct Vulnerability Rating and CVSS against Tenable.")
    ap.add_argument("input", nargs="?", help="Input .csv or .xlsx vuln sheet")
    ap.add_argument("-o", "--output", help="Output path (default: <input>_corrected.<ext>)")
    ap.add_argument("--inplace", action="store_true", help="Overwrite the input file")
    ap.add_argument("--report", help="Write a change report CSV to this path")
    ap.add_argument("--title-col"); ap.add_argument("--rating-col"); ap.add_argument("--cvss-col"); ap.add_argument("--cve-col")
    args = ap.parse_args()

    path = args.input or input("Path to vuln sheet (.csv/.xlsx): ").strip('" ')
    if not path or not os.path.isfile(path):
        sys.exit(f"File not found: {path!r}")

    ext = os.path.splitext(path)[1].lower()
    is_xlsx = ext in (".xlsx", ".xlsm")

    if is_xlsx:
        rows, header, wb, ws = read_xlsx(path)
        fieldnames = header
    else:
        rows, fieldnames = read_csv(path)

    # allow manual column overrides by temporarily injecting them at the front of the key list
    if args.title_col:  TITLE_KEYS.insert(0, args.title_col.lower())
    if args.rating_col: RATING_KEYS.insert(0, args.rating_col.lower())
    if args.cvss_col:   CVSS_KEYS.insert(0, args.cvss_col.lower())
    if args.cve_col:    CVE_KEYS.insert(0, args.cve_col.lower())

    rows, changes, cols = process_rows(rows, fieldnames)

    if args.inplace:
        out = path
    elif args.output:
        out = args.output
    else:
        stem, e = os.path.splitext(path)
        out = f"{stem}_corrected{e}"

    if is_xlsx:
        write_xlsx_inplace(wb, ws, header, rows, cols)
        wb.save(out)
    else:
        write_csv(out, rows, fieldnames)

    print(f"\nDone. {len(changes)} row(s) changed.")
    print(f"Corrected sheet: {out}")
    if changes:
        print("\n  Finding                                             Rating            CVSS")
        for t, orr, nr, oc, nc, src in changes[:40]:
            print(f"  {t[:50]:50} {str(orr):>6} -> {nr:<8} {str(oc):>2} -> {nc:<2}")
        if len(changes) > 40:
            print(f"  ... and {len(changes)-40} more")

    if args.report:
        with open(args.report, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["Vulnerability", "Old_Rating", "New_Rating", "Old_CVSS", "New_CVSS", "Source"])
            w.writerows(changes)
        print(f"Change report: {args.report}")

if __name__ == "__main__":
    main()
